import re
from datetime import date
from pathlib import Path

from monitoring.models import Candidate


HEADER_ALIASES = {
    "sector": {"sector"},
    "cell": {"cell", "celll"},
    "full_name": {"names", "name"},
    "gender": {"gender", "sex"},
    "national_id": {"nationalidpassportnumber", "nationalid", "id"},
    "email": {"email"},
    "phone_number": {"phonenumber", "phone"},
    "qualification": {"qualifications", "qualification"},
    "field_of_study": {"fiedofstudy", "fieldofstudy"},
}


def clean_text(value) -> str:
    return " ".join(str(value or "").split()).strip()


def normalize_header(value) -> str:
    return re.sub(r"[^a-z0-9]", "", clean_text(value).lower())


def normalize_national_id(value) -> str:
    return re.sub(r"\s+", "", clean_text(value))


def find_header_row(rows):
    for row_number, row in enumerate(rows, start=1):
        normalized = {normalize_header(value) for value in row}
        if "names" in normalized and any(alias in normalized for alias in HEADER_ALIASES["national_id"]):
            return row_number, {
                field: next((index for index, value in enumerate(row) if normalize_header(value) in aliases), None)
                for field, aliases in HEADER_ALIASES.items()
            }
    return None, {}


def import_candidates(workbook_path: Path, exam_date: date | None = None, dry_run: bool = False) -> tuple[int, int, list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    imported = skipped = 0
    warnings = []
    try:
        for worksheet in workbook.worksheets:
            buffered_rows = list(worksheet.iter_rows(values_only=True))
            header_row, columns = find_header_row(buffered_rows)
            if header_row is None or columns.get("national_id") is None or columns.get("full_name") is None:
                warnings.append(f"Skipped {worksheet.title}: candidate header not found")
                continue

            district = clean_text(worksheet.title).title()
            for row in buffered_rows[header_row:]:
                values = list(row)
                national_id = normalize_national_id(values[columns["national_id"]])
                full_name = clean_text(values[columns["full_name"]])
                if not national_id or not full_name or not national_id.isdigit():
                    skipped += 1
                    continue

                data = {
                    "full_name": full_name,
                    "district": district,
                    "registered_location": " / ".join(filter(None, [
                        clean_text(values[columns["sector"]]) if columns["sector"] is not None else "",
                        clean_text(values[columns["cell"]]) if columns["cell"] is not None else "",
                    ])),
                    "exam_date": exam_date,
                }
                for field in ("sector", "cell", "gender", "email", "phone_number", "qualification", "field_of_study"):
                    index = columns.get(field)
                    data[field] = clean_text(values[index]) if index is not None else ""

                if not dry_run:
                    Candidate.objects.update_or_create(national_id=national_id, defaults=data)
                imported += 1
    finally:
        workbook.close()
    return imported, skipped, warnings
