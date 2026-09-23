import re
import sys
from pathlib import Path
from datetime import date

import pymysql
from openpyxl import load_workbook
from dotenv import load_dotenv
from decouple import config


HEADER_ALIASES = {
    "sector": {"sector"},
    "cell": {"cell", "celll"},
    "full_name": {"names", "name"},
    "gender": {"gender", "sex"},
    "national_id": {"nationalidpassportnumber", "nationalid", "id"},
    "email": {"email"},
    "phone_number": {"phonenumber", "phone"},
    "qualification": {"qualifications", "qualification"},
    "field_of_study": {"fiedofstudy", "fieldofstudy"},
}


def clean_text(value) -> str:
    return " ".join(str(value or "").split()).strip()


def normalize_header(value) -> str:
    return re.sub(r"[^a-z0-9]", "", clean_text(value).lower())


def normalize_national_id(value) -> str:
    return re.sub(r"\s+", "", clean_text(value))


def find_header_row(rows):
    for row_number, row in enumerate(rows, start=1):
        normalized = {normalize_header(value) for value in row}
        if "names" in normalized and any(alias in normalized for alias in HEADER_ALIASES["national_id"]):
            return row_number, {
                field: next((index for index, value in enumerate(row) if normalize_header(value) in aliases), None)
                for field, aliases in HEADER_ALIASES.items()
            }
    return None, {}


def import_into_mysql(workbook_path: Path, conn_params: dict, exam_date: date | None = None, dry_run: bool = False):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    inserted = skipped = 0
    warnings = []
    conn = None
    try:
        conn = pymysql.connect(**conn_params)
        with conn.cursor() as cur:
            for worksheet in workbook.worksheets:
                buffered_rows = list(worksheet.iter_rows(values_only=True))
                header_row, columns = find_header_row(buffered_rows)
                if header_row is None or columns.get("national_id") is None or columns.get("full_name") is None:
                    warnings.append(f"Skipped {worksheet.title}: candidate header not found")
                    continue

                district = clean_text(worksheet.title).title()
                for row in buffered_rows[header_row:]:
                    values = list(row)
                    national_id = normalize_national_id(values[columns["national_id"]])
                    full_name = clean_text(values[columns["full_name"]])
                    if not national_id or not full_name or not national_id.isdigit():
                        skipped += 1
                        continue

                    data = {
                        "national_id": national_id,
                        "full_name": full_name,
                        "district": district,
                        "registered_location": " / ".join(filter(None, [
                            clean_text(values[columns["sector"]]) if columns["sector"] is not None else "",
                            clean_text(values[columns["cell"]]) if columns["cell"] is not None else "",
                        ])),
                        "exam_date": exam_date.isoformat() if exam_date is not None else None,
                    }
                    for field in ("sector", "cell", "gender", "email", "phone_number", "qualification", "field_of_study"):
                        index = columns.get(field)
                        data[field] = clean_text(values[index]) if index is not None else ""

                    if dry_run:
                        print("DRY:", data)
                    else:
                        sql = (
                            "INSERT INTO monitoring_candidate (national_id, full_name, district, registered_location, exam_date, sector, cell, gender, email, phone_number, qualification, field_of_study, created_at, updated_at) "
                            "VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, NOW(), NOW()) "
                            "ON DUPLICATE KEY UPDATE full_name=VALUES(full_name), district=VALUES(district), registered_location=VALUES(registered_location), exam_date=VALUES(exam_date), sector=VALUES(sector), cell=VALUES(cell), gender=VALUES(gender), email=VALUES(email), phone_number=VALUES(phone_number), qualification=VALUES(qualification), field_of_study=VALUES(field_of_study), updated_at=NOW()"
                        )
                        cur.execute(sql, (
                            data["national_id"], data["full_name"], data["district"], data["registered_location"], data["exam_date"],
                            data["sector"], data["cell"], data["gender"], data["email"], data["phone_number"], data["qualification"], data["field_of_study"],
                        ))
                    inserted += 1
            if not dry_run:
                conn.commit()
    finally:
        workbook.close()
        if conn:
            conn.close()
    return inserted, skipped, warnings


def main(argv):
    if len(argv) < 2:
        print("Usage: seed_candidates_mysql.py <workbook.xlsx> [--exam-date=YYYY-MM-DD] [--dry-run]")
        return 2
    load_dotenv(Path(__file__).resolve().parents[1] / '.env')
    db_name = config('DB_NAME')
    db_user = config('DB_USER', default='root')
    db_password = config('DB_PASSWORD', default='')
    db_host = config('DB_HOST', default='localhost')
    db_port = int(config('DB_PORT', default='3306'))

    conn_params = {
        'host': db_host,
        'port': db_port,
        'user': db_user,
        'password': db_password,
        'database': db_name,
        'cursorclass': pymysql.cursors.Cursor,
        'charset': 'utf8mb4',
    }

    workbook = Path(argv[1])
    exam_date = None
    dry_run = False
    for arg in argv[2:]:
        if arg == '--dry-run':
            dry_run = True
        elif arg.startswith('--exam-date='):
            exam_date = date.fromisoformat(arg.split('=', 1)[1])

    if not workbook.exists():
        print(f"Workbook not found: {workbook}")
        return 2

    try:
        inserted, skipped, warnings = import_into_mysql(workbook, conn_params, exam_date, dry_run)
    except Exception as e:
        print('Error connecting or importing:', e)
        return 3

    for w in warnings:
        print('Warning:', w)
    action = 'would import' if dry_run else 'imported'
    print(f"{action.capitalize()} {inserted} candidates; skipped {skipped} rows.")
    return 0


if __name__ == '__main__':
    raise SystemExit(main(sys.argv))
